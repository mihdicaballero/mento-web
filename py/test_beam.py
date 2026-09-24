"""The glue is plain Python, so it is tested without a browser."""

import base64
import io
import json

import common
import docx
import pytest

import beam

CODES = ["ACI 318-19", "CIRSOC 201-25", "EN 1992-2004"]


def solve(**changes):
    return json.loads(beam.run(json.dumps({**beam.EXAMPLE, **changes})))


@pytest.mark.parametrize("code", CODES)
def test_design_passes_for_every_code(code):
    result = solve(code=code)
    assert result["ok"], result
    assert result["flexure"]["bottom"]["bars"]
    assert result["flexure"]["bottom"]["DCR"] <= 1
    assert result["shear"]["DCR"] <= 1
    assert "svg" not in result  # the page draws the section itself; matplotlib is never loaded there


def test_section_places_every_bar_inside_the_stirrup():
    rebar = {
        "bot": {"n1": 2, "d1": 16, "n2": 1, "d2": 12},
        "top": {"n1": 2, "d1": 10},
        "stirrups": {"n": 1, "d": 8, "s": 15},
    }
    section = solve(mode="check", rebar=rebar)["section"]
    assert section["stirrups"] == {"n": 1, "d": 0.8}
    assert sorted(bar["d"] for bar in section["bars"]) == [1.0, 1.0, 1.2, 1.6, 1.6]
    edge = section["cover"] + section["stirrups"]["d"]
    for bar in section["bars"]:
        assert edge <= bar["x"] - bar["d"] / 2 and bar["x"] + bar["d"] / 2 <= section["width"] - edge + 1e-9
        assert edge <= bar["y"] - bar["d"] / 2 and bar["y"] + bar["d"] / 2 <= section["height"] - edge + 1e-9
    bottom = [bar for bar in section["bars"] if bar["y"] < section["height"] / 2]
    assert [bar["x"] for bar in sorted(bottom, key=lambda bar: bar["x"])][1] == section["width"] / 2


def test_section_moves_bars_that_do_not_fit_to_a_second_row():
    result = solve(width=12, height=30, forces=[{"label": "U", "M_y": 40, "V_z": 50}])
    assert result["ok"], result
    bottom = [bar for bar in result["section"]["bars"] if bar["y"] < result["section"]["height"] / 2]
    assert len({bar["y"] for bar in bottom}) == 2, result["flexure"]["bottom"]["bars"]


def test_a_designed_second_row_stays_in_the_second_row():
    # mento puts 2Ø10 in each row here (n_1 and n_3); listed without the empty n_2 they would
    # read as 4Ø10 in the first row, which does not fit a 12 cm web
    result = solve(width=12, height=30, forces=[{"label": "U", "M_y": 40, "V_z": 50}])
    assert result["layouts"]["bot"] == {"n1": 2, "d1": 10.0, "n3": 2, "d3": 10.0}
    assert result["rows"]["bot"] == ["2Ø10", "2Ø10"]
    assert [label["bars"] for label in result["section"]["labels"]["bot"]] == ["2Ø10", "2Ø10"]
    assert not any(notice["code"] == "spacing" for notice in result["notices"])


def test_check_mode_takes_a_second_row_on_each_face():
    rebar = {
        "bot": {"n1": 3, "d1": 20, "n3": 2, "d3": 16},
        "top": {"n1": 2, "d1": 12, "n3": 2, "d3": 10},
        "stirrups": {"n": 1, "d": 8, "s": 15},
    }
    result = solve(mode="check", rebar=rebar, width=20, height=40)
    assert result["layouts"]["bot"] == {"n1": 3, "d1": 20.0, "n3": 2, "d3": 16.0}
    assert result["rows"] == {"bot": ["3Ø20", "2Ø16"], "top": ["2Ø12", "2Ø10"], "st": ["1eØ8/15 cm"]}
    first, second = (label["y"] for label in result["section"]["labels"]["bot"])
    assert second - first > 2  # the second row sits a bar and a row gap above the first


def test_bars_that_do_not_fit_one_row_fail_as_mento_says():
    rebar = {
        "bot": {"n1": 3, "d1": 20, "n2": 2, "d2": 16},
        "top": {"n1": 2, "d1": 12},
        "stirrups": {"n": 1, "d": 8, "s": 15},
    }
    result = solve(mode="check", rebar=rebar, width=20, height=40, forces=[{"label": "U", "M_y": 60, "V_z": 50}])
    spacing = [notice for notice in result["notices"] if notice["code"] == "spacing"]
    assert spacing and spacing[0]["severity"] == "bad" and spacing[0]["values"]["face"] == "bottom"


DOUBLY = {
    "width": 20,
    "height": 40,
    "forces": [{"label": "U1", "M_y": 150, "V_z": 50}, {"label": "U2", "M_y": -20, "V_z": 40}],
}


def test_a_doubly_reinforced_design_is_not_flagged_above_the_maximum():
    result = solve(**DOUBLY)
    bottom = result["flexure"]["bottom"]
    assert float(bottom["A_s"].split()[0]) > float(bottom["A_s_max"].split()[0])  # past the singly reinforced limit
    assert not any(notice["code"] == "as_above_max" for notice in result["notices"])  # mento: ✅ D.R.
    assert not any(notice["code"] == "as_short" for notice in result["notices"])  # its top covers the compression


def test_top_steel_short_of_the_compression_a_doubly_reinforced_beam_needs_fails():
    rebar = {
        "bot": {"n1": 3, "d1": 20, "n3": 2, "d3": 16},
        "top": {"n1": 2, "d1": 10},
        "stirrups": {"n": 1, "d": 8, "s": 15},
    }
    result = solve(mode="check", rebar=rebar, **DOUBLY)
    (top,) = [
        notice for notice in result["notices"] if notice["code"] == "as_short" and notice["values"]["face"] == "top"
    ]
    assert top["severity"] == "bad"
    # U2's negative moment asks less of the top than U1's positive one does as compression steel
    assert top["values"]["why"] == "compression" and top["values"]["combo"] == "U1"
    assert float(top["values"]["limit"].split()[0]) > 1.99


def test_check_mode_reports_insufficient_rebar():
    rebar = {"bot": {"n1": 2, "d1": 10}, "top": {}, "stirrups": {"n": 1, "d": 6, "s": 20}}
    result = solve(mode="check", rebar=rebar)
    assert result["ok"], result
    assert result["flexure"]["bottom"]["bars"] == "2Ø10"
    assert result["flexure"]["bottom"]["DCR"] > 1


def test_same_diameter_bars_are_counted_together():
    rebar = {"bot": {"n1": 2, "d1": 16, "n2": 1, "d2": 16}, "top": {}, "stirrups": {}}
    assert solve(mode="check", rebar=rebar)["flexure"]["bottom"]["bars"] == "3Ø16"


@pytest.mark.parametrize("field", ["fc", "fy", "width", "height", "cover"])
def test_bad_number_names_the_field(field):
    result = solve(**{field: ""})
    assert result == {"ok": False, "kind": "input", "field": field, "message": "missing"}


def test_no_forces_is_an_input_error():
    assert solve(forces=[{"M_y": 0, "V_z": 0}])["field"] == "forces"


def test_report_is_one_word_file_with_flexure_and_shear():
    (file,) = json.loads(beam.report(json.dumps({**beam.EXAMPLE, "code": "ACI 318-19", "lang": "en", "label": "V1/2"})))
    assert file["name"] == "V1_2 - ACI 318-19.docx"
    document = docx.Document(io.BytesIO(base64.b64decode(file["base64"])))
    text = " ".join(paragraph.text.lower() for paragraph in document.paragraphs)
    assert "flexur" in text and "shear" in text
    assert len(document.tables) >= 8


def test_design_offers_alternatives_with_their_own_dcr():
    result = solve()
    bottom = result["options"]["bot"]
    assert len(bottom) > 1, bottom
    assert bottom[0]["bars"] == result["rebar"]["bot"]  # mento's own pick leads
    assert bottom[0]["dcr"] == result["flexure"]["bottom"]["DCR"]
    assert len({option["signature"] for option in bottom}) == len(bottom)
    for option in bottom:
        assert option["dcr"] > 0 and option["area"].endswith("cm²")
    # more steel in a face can only lower that face's DCR
    by_area = sorted(bottom, key=lambda option: float(option["area"].split()[0]))
    assert [option["dcr"] for option in by_area] == sorted((o["dcr"] for o in by_area), reverse=True)


def test_a_chosen_option_survives_a_recalculation():
    first = solve()
    wanted = first["options"]["bot"][1]
    again = solve(choice={"bot": wanted["signature"]})
    assert again["selected"]["bot"] == 1
    assert again["rebar"]["bot"] == wanted["bars"]
    assert again["ledger"][0]["dcr"] == wanted["dcr"]
    assert again["changed"] == []


def test_a_choice_that_no_longer_exists_falls_back_and_says_so():
    result = solve(choice={"bot": "9x32"})
    assert result["selected"]["bot"] == 0
    assert result["changed"] == ["bot"]


def test_ledger_is_three_rows_capacity_demand_and_governing_combination():
    ledger = solve()["ledger"]
    assert [row["key"] for row in ledger] == ["flexure_bottom", "flexure_top", "shear"]
    bottom, top, shear = ledger
    assert bottom["combo"] == "1.2D+1.6L" and bottom["symbol"] == "ØMn"
    assert bottom["capacity"] > bottom["demand"] == 90.0 and bottom["unit"] == "kNm"
    assert top["combo"] == "0.9D+1.0E"  # the only combination that puts the top in tension
    assert shear["symbol"] == "ØVn" and shear["demand"] == 120.0 and shear["unit"] == "kN"
    assert max(row["dcr"] for row in ledger) == max(
        solve()["flexure"]["bottom"]["DCR"], solve()["shear"]["DCR"], solve()["flexure"]["top"]["DCR"]
    )


def test_a_face_without_demand_has_no_dcr():
    ledger = solve(forces=[{"label": "U", "M_y": 90, "V_z": 120}])["ledger"]
    top = next(row for row in ledger if row["key"] == "flexure_top")
    assert top["dcr"] is None and top["combo"] is None and top["capacity"]


def test_steel_below_the_minimum_is_a_notice_not_a_failure():
    result = solve()
    assert any(notice["code"] == "as_below_min" for notice in result["notices"]), result["notices"]
    assert all(row["dcr"] is None or row["dcr"] <= 1 for row in result["ledger"])


def test_tables_keep_mentos_columns_with_the_units_in_the_header():
    tables = solve()["tables"]
    flexure = tables["flexure"]
    assert flexure["columns"][:3] == ["Label", "Comb.", "Position"]
    assert flexure["columns"][flexure["dcr"]] == "DCR"
    assert flexure["units"][3] == "cm²"
    assert len(flexure["rows"]) == len(beam.EXAMPLE["forces"])
    assert tables["shear"]["columns"][-1] == "DCR"


def test_detailed_results_come_as_tables_too():
    result = solve(lang="en")
    flexure, shear = result["reports"]
    assert flexure["title"] == "BEAM FLEXURE DETAILED RESULTS"
    assert [table["title"] for table in flexure["tables"]][:3] == ["Materials", "Geometry", "Design forces"]
    materials = flexure["tables"][0]
    assert materials["columns"] == ["Variable", "Value", "Unit"]
    assert ["Concrete strength", "fc", "25.0", "MPa"] in materials["rows"]
    check = flexure["tables"][3]
    assert check["columns"] == ["Unit", "Value", "Min.", "Max.", "Ok?"]
    assert all(len(row) == 6 for row in check["rows"])
    assert shear["tables"][-1]["rows"][-1][1:3] == ["DCR", "0.81"]
    assert "BEAM FLEXURE DETAILED RESULTS" in result["detailed"]  # the text stays, for Copiá


def test_a_printed_table_is_cut_where_each_column_starts():
    text = "===== T =====\nName      Variable    Value  Unit\n--------  ----------  -----  ----\nWidth         b          20  cm\nCheck                    ✅\n\n"
    (report,) = common.reports(text)
    assert report["tables"] == [
        {
            "title": "Name",
            "columns": ["Variable", "Value", "Unit"],
            "rows": [["Width", "b", "20", "cm"], ["Check", "", "✅", ""]],
        },
    ]


def test_check_mode_has_no_options():
    rebar = {"bot": {"n1": 3, "d1": 16}, "top": {"n1": 2, "d1": 12}, "stirrups": {"n": 1, "d": 6, "s": 13}}
    result = solve(mode="check", rebar=rebar)
    assert result["options"] == {} and result["selected"] == {}
    assert result["rebar"] == {"bot": "3Ø16", "top": "2Ø12", "st": "1eØ6/13 cm"}


def test_report_joins_flexure_and_shear_without_a_blank_page():
    """The shear part starts on its own page, and nothing empty is left before it to spill over."""
    from docx.oxml.ns import qn

    (file,) = json.loads(beam.report(json.dumps({**beam.EXAMPLE, "lang": "en"})))
    body = docx.Document(io.BytesIO(base64.b64decode(file["base64"]))).element.body
    elements = list(body)
    breaks = [index for index, element in enumerate(elements) if list(element.iter(qn("w:pageBreakBefore")))]
    assert len(breaks) == 1, "one join, one page break"
    assert not [br for br in body.iter(qn("w:br")) if br.get(qn("w:type")) == "page"], "no break paragraphs"
    before = elements[breaks[0] - 1]
    assert before.tag == qn("w:tbl"), "the flexure part ends on its last table, not on an empty paragraph"
    title = "".join(node.text or "" for node in elements[breaks[0]].iter(qn("w:t")))
    assert "shear" in title.lower()
