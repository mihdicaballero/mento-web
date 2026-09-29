"""The three calculators in US customary units: inches, psi, kip and ASTM bar sizes, ACI 318-19."""

import base64
import io
import json
import re

import common
import docx
import pytest

import beam
import slab
import wall

MODULES = [beam, slab, wall]
BARS = re.compile(r"\d+#\d+( \+ \d+#\d+)*")  # 2#6 + 1#4
MESH = re.compile(r"(\d+×)?#\d+ @ \d+(\.\d+)? in")  # #4 @ 12 in, 2×#3 @ 8 in


def solve(module, **changes):
    return json.loads(module.run(json.dumps({**module.EXAMPLE_US, **changes})))


@pytest.mark.parametrize("module", MODULES, ids=lambda module: module.__name__)
def test_the_us_example_passes_between_07_and_09(module):
    result = solve(module)
    assert result["ok"], result
    assert result["units"] == "us" and result["complies"]
    assert not [notice for notice in result["notices"] if notice["severity"] == "bad"]
    governing = max(row["dcr"] or 0 for row in result["ledger"])
    assert 0.7 <= governing <= 0.9, governing


def test_the_us_beam_is_governed_by_its_bottom_flexure():
    """As in the metric example: picking another bottom option moves the verdict."""
    ledger = {row["key"]: row["dcr"] for row in solve(beam)["ledger"]}
    assert max(ledger.values()) == ledger["flexure_bottom"]


def test_the_ledger_speaks_kip_and_kip_ft():
    units = {row["key"]: row["unit"] for module in MODULES for row in solve(module)["ledger"]}
    assert units == {"flexure_bottom": "kip-ft", "flexure_top": "kip-ft", "shear": "kip", "shear_in_plane": "kip"}
    bottom = solve(beam)["ledger"][0]
    assert bottom["demand"] == 72.0 and bottom["capacity"] > bottom["demand"]


def test_a_beam_names_its_bars_by_astm_size():
    result = solve(beam)
    assert BARS.fullmatch(result["rebar"]["bot"]) and BARS.fullmatch(result["rebar"]["top"])
    assert MESH.fullmatch(result["rebar"]["st"])
    for group in ("bot", "top"):
        for option in result["options"][group]:
            assert BARS.fullmatch(option["bars"]) and option["area"].endswith(" in²")
            sizes = [value for key, value in option["layout"].items() if key.startswith("d")]
            assert set(sizes) <= set(common.ASTM_BARS)
    assert all(option["area"].endswith(" in²/ft") for option in result["options"]["st"])
    face = result["flexure"]["bottom"]
    assert face["A_s"].endswith(" in²") and face["M_capacity"].endswith(" kip-ft")
    assert result["shear"]["A_v"].endswith(" in²/ft") and result["shear"]["V_capacity"].endswith(" kip")


@pytest.mark.parametrize("module", [slab, wall], ids=lambda module: module.__name__)
def test_a_mesh_is_a_bar_size_at_whole_inches(module):
    result = solve(module)
    for options in result["options"].values():
        for option in options:
            assert MESH.fullmatch(option["bars"]) and option["area"].endswith(" in²/ft")
            assert option["layout"]["d"] in common.ASTM_BARS
            assert 3 <= option["layout"]["s"] <= 18  # ACI 318-19: never wider than 18 in


def test_check_mode_takes_bar_sizes_and_inches():
    rebar = {"bot": {"n1": 3, "d1": 6}, "top": {"n1": 2, "d1": 5}, "stirrups": {"n": 1, "d": 3, "s": 8}}
    result = solve(beam, mode="check", rebar=rebar)
    assert result["ok"], result
    assert result["rebar"] == {"bot": "3#6", "top": "2#5", "st": "#3 @ 8 in"}
    assert result["flexure"]["bottom"]["A_s"] == "1.33 in²"  # three 0.75 in bars
    assert result["section"]["stirrups"] == {"n": 1, "d": 0.375}
    assert sorted(bar["d"] for bar in result["section"]["bars"]) == [0.625, 0.625, 0.75, 0.75, 0.75]


def test_two_stirrups_are_counted_before_the_bar():
    rebar = {"bot": {"n1": 3, "d1": 6}, "top": {}, "stirrups": {"n": 2, "d": 3, "s": 8}}
    assert solve(beam, mode="check", rebar=rebar)["rebar"]["st"] == "2×#3 @ 8 in"


def test_a_slab_and_a_wall_check_the_mesh_typed():
    result = solve(slab, mode="check", rebar={"bot": {"d": 4, "s": 12}, "top": {}})
    assert result["rebar"]["bot"] == "#4 @ 12 in" and result["area"]["bot"] == "0.20 in²/ft"
    result = solve(wall, mode="check", rebar={"horizontal": {"d": 4, "s": 12}, "vertical": {"d": 4, "s": 12}})
    assert result["rebar"] == {"horizontal": "#4 @ 12 in", "vertical": "#4 @ 12 in"}


def test_a_bar_size_that_does_not_exist_is_an_input_error():
    rebar = {"bot": {"n1": 3, "d1": 12}, "top": {}, "stirrups": {"n": 1, "d": 3, "s": 8}}  # no #12 bar
    assert solve(beam, mode="check", rebar=rebar) == {
        "ok": False,
        "kind": "input",
        "field": "rebar",
        "message": "bar_size",
    }


@pytest.mark.parametrize("module", MODULES, ids=lambda module: module.__name__)
def test_only_aci_is_written_for_us_units(module):
    result = solve(module, code="CIRSOC 201-25")
    assert result == {"ok": False, "kind": "input", "field": "code", "message": "units"}


def test_an_unknown_system_is_an_input_error():
    assert solve(beam, units="imperial")["field"] == "units"


@pytest.mark.parametrize("module", MODULES, ids=lambda module: module.__name__)
def test_si_is_the_system_when_none_is_named(module):
    """Every link made before the units existed carries none: it stays metric."""
    assert "units" not in module.EXAMPLE
    assert json.loads(module.run(json.dumps({**module.EXAMPLE, "units": "si"}))) == json.loads(
        module.run(json.dumps(module.EXAMPLE))
    )


def test_a_chosen_option_survives_a_recalculation():
    wanted = solve(beam)["options"]["bot"][1]
    again = solve(beam, choice={"bot": wanted["signature"]})
    assert again["selected"]["bot"] == 1 and again["rebar"]["bot"] == wanted["bars"] and again["changed"] == []


def test_the_drawing_is_in_inches():
    section = solve(beam)["section"]
    assert (section["width"], section["height"], section["cover"], section["unit"]) == (12, 24, 1.5, "in")
    edge = section["cover"] + section["stirrups"]["d"]
    for bar in section["bars"]:
        assert edge <= bar["x"] - bar["d"] / 2 and bar["x"] + bar["d"] / 2 <= section["width"] - edge + 1e-9
    # a wall's length is typed in feet and drawn in inches, to the same scale as its thickness
    section = solve(wall)["section"]
    assert (section["length"], section["thickness"], section["unit"]) == (120, 8, "in")
    assert all(0 <= bar["x"] <= section["length"] for bar in section["bars"])


def test_the_same_beam_carries_the_same_moment_in_either_system():
    """The units only change how it is written: 12x24 in with 3#6 is 30.48x60.96 cm with 3Ø19.05."""
    rebar_us = {"bot": {"n1": 3, "d1": 6}, "top": {"n1": 2, "d1": 5}, "stirrups": {"n": 1, "d": 3, "s": 8}}
    us = solve(beam, mode="check", rebar=rebar_us)
    rebar_si = {
        "bot": {"n1": 3, "d1": 19.05},
        "top": {"n1": 2, "d1": 15.875},
        "stirrups": {"n": 1, "d": 9.525, "s": 20.32},
    }
    forces = [
        {"label": row["label"], "M_y": row["M_y"] * 1.3558, "V_z": row["V_z"] * 4.4482, "N_x": 0}
        for row in beam.EXAMPLE_US["forces"]
    ]
    si = json.loads(
        beam.run(
            json.dumps(
                {
                    **beam.EXAMPLE,
                    "code": "ACI 318-19",
                    "mode": "check",
                    "fc": 27.579,
                    "fy": 413.685,
                    "width": 30.48,
                    "height": 60.96,
                    "cover": 38.1,
                    "forces": forces,
                    "rebar": rebar_si,
                }
            )
        )
    )
    for row_us, row_si in zip(us["ledger"], si["ledger"], strict=True):
        assert row_us["dcr"] == pytest.approx(row_si["dcr"], rel=0.02), row_us["key"]


def test_the_report_is_written():
    (file,) = json.loads(beam.report(json.dumps(beam.EXAMPLE_US)))
    assert file["name"] == "B-101 - ACI 318-19.docx"
    assert docx.Document(io.BytesIO(base64.b64decode(file["base64"]))).tables


@pytest.mark.xfail(reason="mento 1.3.0 prints an imperial beam in metric; 1.4.0 fixes it", strict=True)
def test_mentos_tables_and_report_are_in_us_units():
    """When this passes, bump MENTO_VERSION and drop the xfail (and common.ASTM_BARS)."""
    result = solve(beam)
    assert "in²" in result["tables"]["flexure"]["units"]
    assert "cm" not in result["detailed"] and "kNm" not in result["detailed"]
    (file,) = json.loads(beam.report(json.dumps(beam.EXAMPLE_US)))
    document = docx.Document(io.BytesIO(base64.b64decode(file["base64"])))
    text = " ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "cm²" not in text and "MPa" not in text
